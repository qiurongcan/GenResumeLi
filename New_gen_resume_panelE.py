# 2025年 生成word版本的简历

# 简历生成分配
import pandas as pd
import os
import random
import subprocess
import glob
from tqdm import tqdm
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
import json
from box import Box
from pprint import pprint

import random

dir_path = r'New_Resume_PanelE'
os.makedirs(dir_path, exist_ok=True)

inform = Box.from_json(filename="information.json")
# pprint(inform)

sheet = 'PanelE_大专生'

df = pd.read_excel(r'简历编码目录.xlsx', index_col=None, sheet_name=sheet)
# print(df.head(3))

# 计算机项目经历


for i in tqdm(range(len(df))):
    
    resume_id = df.iloc[i,0]
    major = df.iloc[i, 1]
    sex = df.iloc[i, 2]
    degree = df.iloc[i, 3]
    level0 = df.iloc[i, 4]
    if degree == "大专":
        degree = "专科"
    
    level = None
    if level0 == "好":
        level = "高质"
    else:
        level = "普通"


    nation = "汉族"

    # 根据性别随机生成名字
    if nation == "汉族":
        name = random.choice(inform.name[sex])
    else:
        name = inform.minority[sex][nation]

    # 出生年月
    birth_year = inform.birth_year[degree] 
    birth_month = random.choice(inform.birth_month)
    
    # 出生地
    birth_place = random.choice(inform.birth_place)
    birth_dist = random.choice(inform.district[birth_place])
    # 现居地
    living_place = random.choice(inform.birth_place)
    district = random.choice(inform.district[living_place])
    
    # 工作实习经历 专科的
    proj = inform.junior_project[major]
    
    # 创建新文档
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = 'Times New Roman'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体') # style，所有文字

    style = doc.styles['Heading 1']
    style.font.name = 'Times New Roman'
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体') # style，所有文字
    style.font.size = Pt(18)

    h0 = doc.add_heading(f"COSER2025 {sheet} {major}", level=1)
    h0.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    h1 = doc.add_heading("1 个人信息", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(f"""
        姓名：{name}
        性别：{sex}
        显示方式：显示
        当前身份：应届毕业生
        出生年月：{birth_year}年{birth_month}月
        现居住城市：{living_place}市{district}
        户口所在地：{birth_place}市{birth_dist}
        政治面貌：共青团员
        手机号码：与注册号码一致
        电子邮箱：写一个自己的邮箱
        微信号：与注册微信一致（空着不写）
    """)

    # 2. 技能熟练程度
    h1 = doc.add_heading("2 技能熟练度", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(proj["professionalSkills"] + '\n设置简历内的显示方式（选择默认选择进度条）')

    # 3.求职状态
    h1 = doc.add_heading("3 求职状态", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph("正在找工作（在“我的”页面点击头像—编辑资料—基本资料，倒数第四行可以设置求职状态为“正在找工作”）")

    # 4.教育经历
    # 这里的专业存在问题 需要修改为：计算机：计算机应用技术； 会计：大数据与会计
    school = random.choice(inform.schools[living_place][level+degree])
    if degree == "专科":
        edu_exp = f"""
    （大专阶段）
    学历：大专-统招
    学校名称：{school}
    所学专业：{inform.junior_major[major]} 
    在校时间：2023.9-2026.6
    """

    h1 = doc.add_heading("4 教育经历", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(edu_exp)


    # 5.在校经历
    h1 = doc.add_heading("5 在校经历-学生职务", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    paragraph = doc.add_paragraph(inform.school_exp_newcoder)

    # 6工作实习经历
    h1 = doc.add_heading("6 工作实习经历", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # 所属行业：{proj['industry']}
    # 拥有技能：{proj["skills"]}
    # 当时月薪：{proj['salary']}
    # （勾选此段经历为实习经历）
    # 对这家公司隐藏我的信息（开启）
    doc.add_paragraph(f"""
    公司名称：{proj['companyName']}
    职位名称：{proj["jobName"]}
    工作类型：实习
    在职时间：{proj['workTime']}
    工作内容：
    {proj['workDescription']}
    """)

    # 7.项目经历
    # 项目经历,复用proj
    h1 = doc.add_heading("7 项目经历", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(proj["experience"])


    # 8.求职意向
    random_city = random.choice(inform.desire_city)

    h1 = doc.add_heading("8 求职意向", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # 期望行业：{inform.tgt_career[major].desire_industry}
    # 工作性质：全职
    doc.add_paragraph(f"""
    期望职位：{inform.tgt_career[major].desire_career}
    求职偏好：空着不写
    工作城市：全国
    薪资要求：{inform.desire_salary[major][level+degree]}
    求职状态：在校-正在找工作
    自定义简历中的位置：设置为“求职意向完整展示在简历内”
    """)

    h1 = doc.add_heading("9 自我评价", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(proj.personAdvantage)

    h1 = doc.add_heading("10 资格证书", level=2)
    h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(inform.certificate)
    h2 = doc.add_heading("退出简历填写，点击“我的简历”-求职状态，填写如下信息", level=3)
    h2.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    doc.add_paragraph(f"""
    1、	我现在：正在找工作
    2、	意向职位：{inform.tgt_career[major].desire_career}
    3、	意向地点：全国
    4、	意向年薪：{inform.desire_salary_year[major][level+degree]}
    """)



    doc.add_paragraph("""
    注：规范简历展示顺序如下（牛客可在“在线简历”右下角“模块管理”自行调整）:
    1、	基本信息
    2、	教育背景
    3、	求职意向
    4、	在校经历
    5、	工作经历
    6、	项目经历
    7、	技能熟练度
    8、	自我评价
    9、	资格证书
    """)



    # # 个人优势 需要修改为专科的
    # personAdvantage = inform.junior_project[major]['personAdvantage']

    # h1 = doc.add_heading("2 个人优势", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(personAdvantage)

    # h1 = doc.add_heading("3 求职状态", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph("在校-正在找工作")

    # random_city = random.choice(inform.desire_city)

    # h1 = doc.add_heading("4 求职意向", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(f"""
    # 期望职位：{inform.tgt_career[major].desire_career}
    # 期望行业：{inform.tgt_career[major].desire_industry}
    # 求职偏好：空着不写
    # 工作城市：北京、上海、广州、西安、{random_city}
    # 薪资要求：{inform.desire_salary[major][level+degree]}
    # 工作性质：全职
    # """)
    # # {inform.desire_salary[major][level+degree]}

    
    
    # h1 = doc.add_heading("5 工作实习经历", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(f"""
    # 职位名称：{proj["jobName"]}
    # 公司名称：{proj['companyName']}
    # 所属行业：{proj['industry']}
    # 在职时间：{proj['workTime']}
    # 工作内容：
    # {proj['workDescription']}
    # 拥有技能：{proj["skills"]}
    # 当时月薪：{proj['salary']}
    # （勾选此段经历为实习经历）
    # 对这家公司隐藏我的信息（开启）
    # """)

    # # 项目经历,复用proj
    # h1 = doc.add_heading("6 项目经历", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(proj["experience"])

    # 

    

    # h1 = doc.add_heading("8 专业技能", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(proj["professionalSkills"])


    # h1 = doc.add_heading("9 资格证书", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # doc.add_paragraph(inform.certificate)


    # h1 = doc.add_heading("10 学生干部经历", level=2)
    # h1.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    # paragraph = doc.add_paragraph(inform.school_exp)

    # 设置字体
    # for paragraph in doc.paragraphs:
    # for run in paragraph.runs:
    #     # run.font.name = u'宋体'
    #     run.font.size = Pt(14)

    # if major == "STEM":
    #     major_id = 1
    # else:
    #     major_id = 2

    doc.save(f"{dir_path}/{resume_id}.docx")
    # break
    # if i == 10:
        # break








